# -*- coding: utf-8 -*-
"""
Generate high-quality PDF and DOCX downloadable documents for all 4 LexAI team member case studies
and the master portfolio overview.
"""

import os
import re
import zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas

# Define Palette
NAVY = colors.HexColor("#0F2D5F")
ORANGE = colors.HexColor("#FF6A0D")
DARK_GRAY = colors.HexColor("#2D3748")
LIGHT_BG = colors.HexColor("#F8FAFC")
BORDER_COLOR = colors.HexColor("#E2E8F0")
GREEN = colors.HexColor("#10B981")
BLUE = colors.HexColor("#2563EB")
WHITE = colors.white

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        
        # Header (Pages > 1)
        if self._pageNumber > 1:
            self.drawString(54, 755, "LexAI — Natural Language Processing Legal Assistant Case Study")
            self.setStrokeColor(BORDER_COLOR)
            self.setLineWidth(0.5)
            self.line(54, 747, 558, 747)
            
        # Footer
        page_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(558, 35, page_text)
        self.drawString(54, 35, "Department of Computer Science & Engineering | Confidential Academic Deliverable")
        self.setStrokeColor(BORDER_COLOR)
        self.setLineWidth(0.5)
        self.line(54, 45, 558, 45)
        
        self.restoreState()


def parse_markdown_to_flowables(md_content, styles):
    flowables = []
    lines = md_content.split('\n')
    
    in_table = False
    table_rows = []
    in_code = False
    code_lines = []
    
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        
        # Check code block
        if line.startswith('```'):
            if in_code:
                # End code block
                in_code = False
                code_text = "<br/>".join(code_lines).replace(" ", "&nbsp;")
                p = Paragraph(f"<font face='Courier' size=8 color='#1E293B'>{code_text}</font>", styles['CodeBlock'])
                t = Table([[p]], colWidths=[500])
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#F1F5F9")),
                    ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
                    ('TOPPADDING', (0,0), (-1,-1), 6),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 6),
                    ('LEFTPADDING', (0,0), (-1,-1), 8),
                    ('RIGHTPADDING', (0,0), (-1,-1), 8),
                ]))
                flowables.append(Spacer(1, 4))
                flowables.append(t)
                flowables.append(Spacer(1, 6))
                code_lines = []
            else:
                in_code = True
                code_lines = []
            i += 1
            continue
            
        if in_code:
            clean_code_line = line.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            code_lines.append(clean_code_line)
            i += 1
            continue
            
        # Check table
        if line.strip().startswith('|') and line.strip().endswith('|'):
            # Table line
            raw_cells = [c.strip() for c in line.strip()[1:-1].split('|')]
            # Check if separator row like |:---|:---|
            if all(re.match(r'^:?-+:?$', c) for c in raw_cells):
                i += 1
                continue
            table_rows.append(raw_cells)
            in_table = True
            i += 1
            continue
        else:
            if in_table:
                # Flush table
                in_table = False
                if table_rows:
                    num_cols = max(len(r) for r in table_rows)
                    # Normalize rows
                    norm_rows = []
                    for r_idx, r in enumerate(table_rows):
                        row_cells = []
                        is_header = (r_idx == 0)
                        for c_idx in range(num_cols):
                            val = r[c_idx] if c_idx < len(r) else ""
                            # Format formatting in cells
                            val_html = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', val)
                            val_html = re.sub(r'\*(.*?)\*', r'<i>\1</i>', val_html)
                            val_html = re.sub(r'`(.*?)`', r"<font face='Courier'>\1</font>", val_html)
                            val_html = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', val_html)
                            style = styles['TableHeader'] if is_header else styles['TableCell']
                            p = Paragraph(val_html, style)
                            row_cells.append(p)
                        norm_rows.append(row_cells)
                    
                    # Col widths based on count
                    col_w = 504 / num_cols
                    col_widths = [col_w] * num_cols
                    # Custom adjustment if 2, 4 or 5 columns
                    if num_cols == 2:
                        col_widths = [160, 344]
                    elif num_cols == 5:
                        col_widths = [70, 70, 110, 154, 100]
                    elif num_cols == 4:
                        col_widths = [110, 110, 140, 144]
                    elif num_cols == 6:
                        col_widths = [120, 60, 60, 64, 50, 150]
                        
                    t = Table(norm_rows, colWidths=col_widths)
                    t.setStyle(TableStyle([
                        ('BACKGROUND', (0,0), (-1,0), NAVY),
                        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
                        ('TOPPADDING', (0,0), (-1,-1), 4),
                        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
                        ('LEFTPADDING', (0,0), (-1,-1), 5),
                        ('RIGHTPADDING', (0,0), (-1,-1), 5),
                        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
                    ]))
                    flowables.append(Spacer(1, 4))
                    flowables.append(t)
                    flowables.append(Spacer(1, 8))
                table_rows = []

        # Empty lines
        if not line.strip():
            flowables.append(Spacer(1, 4))
            i += 1
            continue

        # Horizontal rule
        if line.strip() in ['---', '***', '___']:
            flowables.append(Spacer(1, 4))
            flowables.append(HRFlowable(width="100%", thickness=1, color=BORDER_COLOR, spaceBefore=4, spaceAfter=8))
            i += 1
            continue

        # Headers
        if line.startswith('# '):
            text = line[2:].strip()
            p = Paragraph(f"<b>{text}</b>", styles['MainTitle'])
            flowables.append(p)
            flowables.append(Spacer(1, 4))
        elif line.startswith('## '):
            text = line[3:].strip()
            p = Paragraph(f"<b>{text}</b>", styles['Heading1Custom'])
            flowables.append(Spacer(1, 8))
            flowables.append(p)
            flowables.append(Spacer(1, 4))
        elif line.startswith('### '):
            text = line[4:].strip()
            p = Paragraph(f"<b>{text}</b>", styles['Heading2Custom'])
            flowables.append(Spacer(1, 6))
            flowables.append(p)
            flowables.append(Spacer(1, 3))
        elif line.startswith('#### '):
            text = line[5:].strip()
            p = Paragraph(f"<b>{text}</b>", styles['Heading3Custom'])
            flowables.append(Spacer(1, 4))
            flowables.append(p)
            flowables.append(Spacer(1, 2))
        elif line.startswith('- ') or line.startswith('* ') or re.match(r'^\d+\.\s', line):
            # Bullet or numbered item
            prefix = "&bull;&nbsp;&nbsp;"
            if re.match(r'^\d+\.\s', line):
                match = re.match(r'^(\d+\.)\s', line)
                prefix = f"<b>{match.group(1)}</b>&nbsp;&nbsp;"
                body = line[match.end():]
            else:
                body = line[2:]
            
            body_html = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', body)
            body_html = re.sub(r'\*(.*?)\*', r'<i>\1</i>', body_html)
            body_html = re.sub(r'`(.*?)`', r"<font face='Courier'>\1</font>", body_html)
            body_html = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', body_html)
            p = Paragraph(f"{prefix}{body_html}", styles['BulletCustom'])
            flowables.append(p)
        else:
            # Normal paragraph
            body_html = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', line)
            body_html = re.sub(r'\*(.*?)\*', r'<i>\1</i>', body_html)
            body_html = re.sub(r'`(.*?)`', r"<font face='Courier'>\1</font>", body_html)
            body_html = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', body_html)
            p = Paragraph(body_html, styles['BodyCustom'])
            flowables.append(p)
            
        i += 1
        
    return flowables


def build_pdf_from_md(md_path, pdf_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )
    
    styles = getSampleStyleSheet()
    
    styles.add(ParagraphStyle(
        name='MainTitle',
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=18,
        textColor=NAVY,
        alignment=0,
        spaceAfter=4
    ))
    styles.add(ParagraphStyle(
        name='Heading1Custom',
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=15,
        textColor=NAVY,
        spaceBefore=8,
        spaceAfter=4,
        keepWithNext=True
    ))
    styles.add(ParagraphStyle(
        name='Heading2Custom',
        fontName='Helvetica-Bold',
        fontSize=10.5,
        leading=13,
        textColor=ORANGE,
        spaceBefore=6,
        spaceAfter=3,
        keepWithNext=True
    ))
    styles.add(ParagraphStyle(
        name='Heading3Custom',
        fontName='Helvetica-Bold',
        fontSize=9.5,
        leading=12,
        textColor=DARK_GRAY,
        spaceBefore=4,
        spaceAfter=2,
        keepWithNext=True
    ))
    styles.add(ParagraphStyle(
        name='BodyCustom',
        fontName='Helvetica',
        fontSize=8.5,
        leading=11.5,
        textColor=DARK_GRAY,
        spaceAfter=4
    ))
    styles.add(ParagraphStyle(
        name='BulletCustom',
        fontName='Helvetica',
        fontSize=8.5,
        leading=11.5,
        textColor=DARK_GRAY,
        leftIndent=14,
        spaceAfter=2.5
    ))
    styles.add(ParagraphStyle(
        name='TableHeader',
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=WHITE,
        alignment=0
    ))
    styles.add(ParagraphStyle(
        name='TableCell',
        fontName='Helvetica',
        fontSize=7.8,
        leading=9.8,
        textColor=DARK_GRAY,
        alignment=0
    ))
    styles.add(ParagraphStyle(
        name='CodeBlock',
        fontName='Courier',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#1E293B")
    ))

    flowables = parse_markdown_to_flowables(md_text, styles)
    doc.build(flowables, canvasmaker=NumberedCanvas)
    print(f"Generated PDF: {pdf_path}")


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def build_docx_from_md(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    doc = Document()
    
    # Set page margins to 0.75 in
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    lines = md_text.split('\n')
    in_table = False
    table_rows = []
    in_code = False
    code_lines = []

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        # Code block
        if line.startswith('```'):
            if in_code:
                in_code = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.2)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run('\n'.join(code_lines))
                run.font.name = 'Consolas'
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(30, 41, 59)
                code_lines = []
            else:
                in_code = True
                code_lines = []
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # Table detection
        if line.strip().startswith('|') and line.strip().endswith('|'):
            raw_cells = [c.strip() for c in line.strip()[1:-1].split('|')]
            if all(re.match(r'^:?-+:?$', c) for c in raw_cells):
                i += 1
                continue
            table_rows.append(raw_cells)
            in_table = True
            i += 1
            continue
        else:
            if in_table:
                in_table = False
                if table_rows:
                    num_rows = len(table_rows)
                    num_cols = max(len(r) for r in table_rows)
                    tbl = doc.add_table(rows=num_rows, cols=num_cols)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    tbl.autofit = True

                    for r_idx, row_data in enumerate(table_rows):
                        row = tbl.rows[r_idx]
                        is_header = (r_idx == 0)
                        for c_idx in range(num_cols):
                            cell = row.cells[c_idx]
                            val = row_data[c_idx] if c_idx < len(row_data) else ""
                            # Strip formatting for text
                            val_clean = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', val)
                            cell.text = ""
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            run = p.add_run(val_clean.replace('**', '').replace('*', ''))
                            run.font.name = 'Times New Roman'
                            run.font.size = Pt(8.5)

                            if is_header:
                                run.font.bold = True
                                run.font.color.rgb = RGBColor(255, 255, 255)
                                set_cell_background(cell, "0F2D5F")
                            else:
                                if r_idx % 2 == 1:
                                    set_cell_background(cell, "FFFFFF")
                                else:
                                    set_cell_background(cell, "F8FAFC")
                            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)

                    doc.add_paragraph().paragraph_format.space_after = Pt(4)
                table_rows = []

        if not line.strip():
            i += 1
            continue

        if line.strip() in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E1"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            i += 1
            continue

        if line.startswith('# '):
            text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 45, 95)
        elif line.startswith('## '):
            text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 45, 95)
        elif line.startswith('### '):
            text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 106, 13)
        elif line.startswith('- ') or line.startswith('* ') or re.match(r'^\d+\.\s', line):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            body = re.sub(r'^[-*]\s+|\d+\.\s+', '', line)
            parts = re.split(r'(\*\*.*?\*\*)', body)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                else:
                    run = p.add_run(part)
                run.font.name = 'Times New Roman'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(45, 55, 72)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                else:
                    run = p.add_run(part)
                run.font.name = 'Times New Roman'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(45, 55, 72)

        i += 1

    doc.save(docx_path)
    print(f"Generated DOCX: {docx_path}")


def main():
    base_dir = r"c:\Users\venga\OneDrive\Desktop\AI legal\case_studies"
    downloads_dir = os.path.join(base_dir, "downloads")
    desktop_dir = r"C:\Users\venga\OneDrive\Desktop\AI_Legal_Case_Studies_Download"
    
    os.makedirs(downloads_dir, exist_ok=True)
    os.makedirs(desktop_dir, exist_ok=True)

    files_to_convert = [
        ("CASE_STUDY_1_NLP_MACHINE_LEARNING_ENGINEER_BRUNDA.md", "Case_Study_1_V_Brunda_NLP_Engineer"),
        ("CASE_STUDY_2_LEGAL_RISK_AND_EXTRACTION_SPECIALIST_KRISHNA.md", "Case_Study_2_G_Krishna_Legal_Risk_Specialist"),
        ("CASE_STUDY_3_BACKEND_RAG_INTEGRATION_ENGINEER_JAHNAVI.md", "Case_Study_3_G_Jahnavi_Backend_RAG_Engineer"),
        ("CASE_STUDY_4_FRONTEND_LEGAL_UX_ENGINEER_MAMATHA.md", "Case_Study_4_K_Mamatha_Frontend_UX_Engineer"),
        ("CASE_STUDY_PORTFOLIO_OVERVIEW.md", "Case_Study_Portfolio_Overview_All_Members")
    ]

    all_generated_files = []

    for md_fname, out_basename in files_to_convert:
        md_path = os.path.join(base_dir, md_fname)
        
        pdf_path = os.path.join(downloads_dir, f"{out_basename}.pdf")
        docx_path = os.path.join(downloads_dir, f"{out_basename}.docx")
        
        # Build PDF
        build_pdf_from_md(md_path, pdf_path)
        # Build DOCX
        build_docx_from_md(md_path, docx_path)
        
        # Also copy to desktop directory
        desktop_pdf = os.path.join(desktop_dir, f"{out_basename}.pdf")
        desktop_docx = os.path.join(desktop_dir, f"{out_basename}.docx")
        
        import shutil
        shutil.copy2(pdf_path, desktop_pdf)
        shutil.copy2(docx_path, desktop_docx)
        
        all_generated_files.extend([desktop_pdf, desktop_docx])

    # Create a Master ZIP archive containing everything
    zip_path_desktop = r"C:\Users\venga\OneDrive\Desktop\AI_Legal_Case_Studies_All_4_Members.zip"
    zip_path_local = os.path.join(downloads_dir, "AI_Legal_Case_Studies_All_4_Members.zip")
    
    with zipfile.ZipFile(zip_path_desktop, 'w', zipfile.ZIP_DEFLATED) as zipf:
        for f in all_generated_files:
            zipf.write(f, arcname=os.path.basename(f))
        for md_fname, _ in files_to_convert:
            zipf.write(os.path.join(base_dir, md_fname), arcname=md_fname)

    import shutil
    shutil.copy2(zip_path_desktop, zip_path_local)
    print(f"\nCreated Master ZIP Archive: {zip_path_desktop}")
    print("All downloads successfully generated!")

if __name__ == "__main__":
    main()
